import io
import zipfile

from docx import Document
from sqlalchemy import select

from frota import db as dbmod
from frota.models import Comodato, Veiculo
from frota.services import colagem

from .conftest import docx_modelo

PLANILHA = (
    "Cliente\tCNPJ\tPlaca\tData Saída\tDevolução\tCentro de custo\n"
    "Acme Eventos Ltda\t12.345.678/0001-90\tabc1d23\t01/11/2026\t05/11/2026\tCC-10\n"
    "Beta Promoções\t\tXYZ9K88\t02/11/2026\t07/11/2026\tCC-20\n"
    "Gama\t\tNAO0000\t02/11/2026\t\tCC-30\n"
)


def _frota(client):
    for placa, modelo in (("ABC1D23", "Onix"), ("XYZ9K88", "HB20")):
        client.post("/veiculos", data={"placa": placa, "marca": "X", "modelo": modelo})


def test_parse_colagem_do_excel():
    cols, linhas = colagem.ler_tabela(PLANILHA)
    assert {"cliente_nome", "placa", "data_saida", "data_prevista_devolucao"} <= set(cols)
    assert linhas[0].valores["placa"] == "ABC1D23"
    assert linhas[0].valores["cliente_documento"] == "12345678000190"
    assert linhas[0].extras == {"centro_de_custo": "CC-10"}


def test_colagem_com_ponto_e_virgula_e_data_invalida():
    _, linhas = colagem.ler_tabela("cliente;placa;data_saida\nAcme;aaa1a11;31/02/2026")
    assert linhas[0].valores["placa"] == "AAA1A11"
    assert any("data_saida" in e for e in linhas[0].erros)


def test_lote_completo_e_documento_word(client):
    _frota(client)
    r = client.post("/modelos", data={"nome": "Comodato padrao"},
                    files={"arquivo": ("m.docx", docx_modelo())})
    assert r.status_code == 303

    prev = client.post("/comodatos/lote", data={"texto": PLANILHA})
    assert "2 ok" in prev.text and "1 com erro" in prev.text and "NAO0000 não cadastrada" in prev.text

    r = client.post("/comodatos/lote", data={"texto": PLANILHA, "confirmar": "1"})
    assert "2 comodato(s) criado(s)" in r.text

    with dbmod.SessionLocal() as db:
        coms = db.scalars(select(Comodato).order_by(Comodato.numero)).all()
        assert [c.veiculo.placa for c in coms] == ["ABC1D23", "XYZ9K88"]
        assert all(c.veiculo.status == "em_comodato" for c in coms)
        cid = coms[0].id
        from frota.models import ModeloDocumento
        modelo_id = db.scalar(select(ModeloDocumento.id))
        ids = ",".join(c.id for c in coms)

    r = client.get(f"/comodatos/{cid}/documento/{modelo_id}")
    assert r.status_code == 200
    texto = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    assert "COMODATO Nº 0001/2026" in texto
    assert "Comodatário: Acme Eventos Ltda, CNPJ 12.345.678/0001-90" in texto  # campo partido em 2 runs
    assert "placa ABC1D23" in texto and "Retirada em 01/11/2026" in texto
    assert "Centro de custo: CC-10" in texto  # coluna extra da planilha
    assert "{{" not in texto

    z = zipfile.ZipFile(io.BytesIO(client.get(f"/comodatos/zip/{modelo_id}?ids={ids}").content))
    assert len(z.namelist()) == 2


def test_carro_ocupado_nao_vira_segundo_comodato(client):
    _frota(client)
    client.post("/comodatos/lote", data={"texto": PLANILHA, "confirmar": "1"})
    prev = client.post("/comodatos/lote", data={"texto": PLANILHA})
    assert "ABC1D23 não está disponível" in prev.text


def test_encerrar_libera_o_carro_e_registra_historico(client):
    _frota(client)
    client.post("/comodatos/lote", data={"texto": PLANILHA, "confirmar": "1"})
    with dbmod.SessionLocal() as db:
        com = db.scalar(select(Comodato).order_by(Comodato.numero))
        cid, vid = com.id, com.veiculo_id
    assert client.post(f"/comodatos/{cid}/encerrar").status_code == 303
    with dbmod.SessionLocal() as db:
        v = db.get(Veiculo, vid)
        assert v.status == "disponivel"
        assert [e.tipo for e in v.eventos] == ["comodato_aberto", "comodato_encerrado"]


def test_modelo_com_campo_desconhecido_falha_claro(client):
    _frota(client)
    d = Document(); d.add_paragraph("{{ campo_que_nao_existe }}")
    b = io.BytesIO(); d.save(b)
    client.post("/modelos", data={"nome": "Ruim"}, files={"arquivo": ("r.docx", b.getvalue())})
    client.post("/comodatos/lote", data={"texto": PLANILHA, "confirmar": "1"})
    with dbmod.SessionLocal() as db:
        from frota.models import ModeloDocumento
        cid = db.scalar(select(Comodato.id)); mid = db.scalar(select(ModeloDocumento.id))
    r = client.get(f"/comodatos/{cid}/documento/{mid}")
    assert r.status_code == 422 and "campo_que_nao_existe" in r.text


def test_paginas_abrem(client):
    for rota in ("/veiculos", "/clientes", "/comodatos", "/comodatos/lote", "/modelos"):
        assert client.get(rota).status_code == 200, rota
