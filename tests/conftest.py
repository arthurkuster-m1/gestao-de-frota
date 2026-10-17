import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from frota import db as dbmod


def docx_modelo() -> bytes:
    """Word de teste. O 2º parágrafo tem o campo partido em dois 'runs', como o Word faz."""
    d = Document()
    d.add_paragraph("COMODATO Nº {{ numero }}/{{ ano }}")
    p = d.add_paragraph()
    p.add_run("Comodatário: {{ cliente_")
    p.add_run("nome }}, CNPJ {{ cliente_documento }}")
    d.add_paragraph("Veículo {{ veiculo_marca }} {{ veiculo_modelo }}, placa {{ veiculo_placa }}.")
    d.add_paragraph("Retirada em {{ data_saida }}; devolução até {{ data_prevista_devolucao }}. Centro de custo: {{ centro_de_custo }}")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("FROTA_DATA_DIR", str(tmp_path))
    dbmod.init_db()
    from frota.main import app

    with TestClient(app, follow_redirects=False) as c:
        yield c
